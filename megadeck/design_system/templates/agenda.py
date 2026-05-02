"""Agenda slide — numbered topic list with descriptions."""
from __future__ import annotations

from pptx.enum.text import MSO_ANCHOR
from pptx.slide import Slide
from pptx.util import Inches, Pt

from megadeck.core.schemas import AgendaSlide
from megadeck.design_system.primitives import (
    add_corner_dotgrid,
    add_eyebrow,
    add_page_chrome,
    add_round_rect,
    add_text,
    set_slide_bg,
)
from megadeck.design_system.tokens import Theme


def render_agenda(
    slide: Slide,
    data: AgendaSlide,
    theme: Theme,
    *,
    page_n: int,
    page_total: int,
    section_label: str | None = None,
) -> None:
    set_slide_bg(slide, color=theme.bg, theme=theme)
    add_corner_dotgrid(slide, theme)
    add_eyebrow(slide, text="AGENDA", theme=theme)

    LEFT = theme.left_margin_in
    CONTENT_W = theme.content_width_in

    add_text(
        slide,
        left=LEFT, top=1.20,
        width=CONTENT_W, height=0.95,
        text=data.title,
        font=theme.font_display,
        size_pt=theme.type_scale.h2 + 4,
        color=theme.title,
        bold=True,
    )
    body_top = 2.40
    bottom = 6.95
    available = bottom - body_top

    n = len(data.items)
    gap = 0.18 if n <= 5 else 0.12
    item_h = max(0.55, (available - (n - 1) * gap) / n)

    for i, agenda in enumerate(data.items):
        y = body_top + i * (item_h + gap)
        # Soft surface card
        add_round_rect(
            slide,
            left=LEFT, top=y, width=CONTENT_W, height=item_h,
            fill=theme.surface, line=theme.hairline, line_w=0.5,
            adjust=0.10,
        )
        # Number column (left)
        add_text(
            slide,
            left=LEFT + 0.30, top=y, width=1.10, height=item_h,
            text=agenda.number,
            font=theme.font_display,
            size_pt=theme.type_scale.h3,
            color=theme.accent,
            bold=True,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        # Title + description (rich runs)
        tb = slide.shapes.add_textbox(
            Inches(LEFT + 1.50), Inches(y),
            Inches(CONTENT_W - 1.70), Inches(item_h),
        )
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.04)
        tf.margin_top = Inches(0.02)
        try:
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        except Exception:
            pass
        p = tf.paragraphs[0]
        try:
            p.line_spacing = 1.15
        except Exception:
            pass
        r1 = p.add_run()
        r1.text = agenda.title
        r1.font.name = theme.font_display
        r1.font.size = Pt(theme.type_scale.h4 - 2)
        r1.font.bold = True
        r1.font.color.rgb = theme.title
        if agenda.description:
            r2 = p.add_run()
            r2.text = "    " + agenda.description
            r2.font.name = theme.font_body
            r2.font.size = Pt(theme.type_scale.micro + 1)
            r2.font.color.rgb = theme.body

    add_page_chrome(
        slide, theme=theme,
        page_n=page_n, page_total=page_total,
        section_label=section_label,
    )
