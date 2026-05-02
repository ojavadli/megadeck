"""Two-column comparison slide — left vs right with thin accent bars."""
from __future__ import annotations

from pptx.enum.text import MSO_ANCHOR
from pptx.slide import Slide
from pptx.util import Inches, Pt

from megadeck.core.schemas import TwoColumnSlide, BulletItem
from megadeck.design_system.primitives import (
    add_corner_dotgrid,
    add_eyebrow,
    add_oval,
    add_page_chrome,
    add_text,
    add_v_line,
    set_slide_bg,
)
from megadeck.design_system.tokens import Theme


def _build_column(
    slide: Slide,
    *,
    x: float,
    body_top: float,
    available_h: float,
    col_w: float,
    col_title: str,
    items: list[BulletItem],
    theme: Theme,
) -> None:
    # Column header — small dot + title
    add_oval(
        slide,
        left=x, top=body_top + 0.13,
        size=0.10,
        fill=theme.accent,
    )
    add_text(
        slide,
        left=x + 0.22, top=body_top,
        width=col_w - 0.22, height=0.45,
        text=col_title,
        font=theme.font_display,
        size_pt=theme.type_scale.h4 - 4,
        color=theme.title,
        bold=True,
    )
    list_top = body_top + 0.65
    list_h = available_h - 0.65
    n = len(items)
    if n <= 3:
        head_pt, tail_pt, gap = 17, 15, 0.20
    elif n == 4:
        head_pt, tail_pt, gap = 16, 14, 0.16
    else:
        head_pt, tail_pt, gap = 14, 13, 0.13
    item_h = max(0.55, (list_h - (n - 1) * gap) / n)

    for i, item in enumerate(items):
        y = list_top + i * (item_h + gap)
        bar_pad = 0.05
        add_v_line(
            slide,
            left=x, top=y + bar_pad,
            height=item_h - 2 * bar_pad,
            color=theme.accent,
            width=0.03,
        )
        tb = slide.shapes.add_textbox(
            Inches(x + 0.22), Inches(y),
            Inches(col_w - 0.30), Inches(item_h),
        )
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.04)
        tf.margin_top = Inches(0.02)
        tf.margin_bottom = Inches(0.02)
        try:
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        except Exception:
            pass
        p = tf.paragraphs[0]
        try:
            p.line_spacing = 1.15
        except Exception:
            pass
        r1 = p.add_run()
        r1.text = item.head
        r1.font.name = theme.font_display
        r1.font.size = Pt(head_pt)
        r1.font.bold = True
        r1.font.color.rgb = theme.title
        if item.tail:
            r2 = p.add_run()
            r2.text = "  " + item.tail
            r2.font.name = theme.font_body
            r2.font.size = Pt(tail_pt)
            r2.font.color.rgb = theme.body


def render_two_column(
    slide: Slide,
    data: TwoColumnSlide,
    theme: Theme,
    *,
    page_n: int,
    page_total: int,
    section_label: str | None = None,
) -> None:
    set_slide_bg(slide, color=theme.bg, theme=theme)
    add_corner_dotgrid(slide, theme)
    add_eyebrow(slide, text=data.eyebrow.upper(), theme=theme)

    LEFT = theme.left_margin_in
    CONTENT_W = theme.content_width_in

    add_text(
        slide,
        left=LEFT, top=1.20,
        width=CONTENT_W, height=0.85,
        text=data.title,
        font=theme.font_display,
        size_pt=theme.type_scale.h2 - 2,
        color=theme.title,
        bold=True,
    )
    title_offset = 1.20 + 0.95
    if data.subtitle:
        add_text(
            slide,
            left=LEFT, top=title_offset,
            width=CONTENT_W, height=0.40,
            text=data.subtitle,
            font=theme.font_body,
            size_pt=theme.type_scale.micro,
            color=theme.muted,
        )
        body_top = title_offset + 0.55
    else:
        body_top = title_offset + 0.10

    bottom = 6.95
    available = bottom - body_top
    col_gap = 0.50
    col_w = (CONTENT_W - col_gap) / 2

    _build_column(
        slide, x=LEFT, body_top=body_top, available_h=available,
        col_w=col_w, col_title=data.left_title,
        items=data.left_items, theme=theme,
    )
    _build_column(
        slide, x=LEFT + col_w + col_gap, body_top=body_top, available_h=available,
        col_w=col_w, col_title=data.right_title,
        items=data.right_items, theme=theme,
    )
    add_page_chrome(
        slide, theme=theme,
        page_n=page_n, page_total=page_total,
        section_label=section_label,
    )
