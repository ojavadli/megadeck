"""Hero statement slide — one massive line, supporting copy below."""
from __future__ import annotations

from pptx.slide import Slide
from pptx.util import Inches, Pt

from megadeck.core.schemas import HeroStatementSlide
from megadeck.design_system.primitives import (
    add_corner_dotgrid,
    add_eyebrow,
    add_page_chrome,
    add_rect,
    add_text,
    fit_title,
    measure_title_height,
    set_slide_bg,
)
from megadeck.design_system.tokens import Theme


def render_hero_statement(
    slide: Slide,
    data: HeroStatementSlide,
    theme: Theme,
    *,
    page_n: int,
    page_total: int,
    section_label: str | None = None,
) -> None:
    set_slide_bg(slide, color=theme.bg, theme=theme)
    add_corner_dotgrid(slide, theme)
    add_eyebrow(slide, text=data.eyebrow.upper(), theme=theme, top=0.95)

    LEFT = theme.left_margin_in
    CONTENT_W = theme.content_width_in

    statement_pt = fit_title(
        data.statement, max_pt=theme.type_scale.hero, min_pt=36, width_in=CONTENT_W,
    )
    statement_h = measure_title_height(
        data.statement, size_pt=statement_pt, width_in=CONTENT_W, line_spacing=1.0,
    )
    statement_top = 2.10
    statement_tb = add_text(
        slide,
        left=LEFT, top=statement_top,
        width=CONTENT_W, height=statement_h,
        text=data.statement,
        font=theme.font_display,
        size_pt=statement_pt,
        color=theme.title,
        bold=True,
        line_spacing=1.0,
    )
    if theme.accent_glow and theme.is_dark:
        try:
            from megadeck.design_system.effects import apply_glow
            apply_glow(
                statement_tb,
                color=theme.accent,
                radius_pt=18.0,
                alpha_pct=35,
            )
        except Exception:
            pass

    dash_top = statement_top + statement_h + 0.20
    dash = add_rect(
        slide,
        left=LEFT, top=dash_top,
        width=0.55, height=0.06,
        fill=theme.accent,
    )
    if theme.accent_glow:
        try:
            from megadeck.design_system.effects import apply_glow
            apply_glow(dash, color=theme.accent, radius_pt=12.0, alpha_pct=70)
        except Exception:
            pass

    if data.supports:
        bottom = 7.00
        sup_top = dash_top + 0.30
        available = max(0.5, bottom - sup_top)
        n = len(data.supports)
        # Auto-shrink supports if multi-line: each support that wraps to 2 lines
        # would otherwise overlap a fixed-height neighbour. Use a single
        # multi-paragraph textbox so PowerPoint flows them with proper line gaps.
        body_pt = theme.type_scale.body_large
        # Estimate total wanted height to choose a font size that actually fits.
        for try_pt in (body_pt, body_pt - 2, body_pt - 4, body_pt - 6, body_pt - 8):
            chars_per_line = max(20, int(CONTENT_W * 72.0 / (try_pt * 0.50)))
            wanted_lines = sum(
                max(1, (len(s) + chars_per_line - 1) // chars_per_line)
                for s in data.supports
            )
            line_h = (try_pt * 1.30) / 72.0
            para_gap = 0.18
            wanted_h = wanted_lines * line_h + (n - 1) * para_gap
            if wanted_h <= available:
                body_pt = try_pt
                break
        tb = slide.shapes.add_textbox(
            Inches(LEFT), Inches(sup_top), Inches(CONTENT_W), Inches(available),
        )
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.0)
        tf.margin_top = Inches(0.0)
        tf.margin_right = Inches(0.0)
        tf.margin_bottom = Inches(0.0)
        for i, line in enumerate(data.supports):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            try:
                p.line_spacing = 1.30
                if i > 0:
                    p.space_before = Pt(8)
            except Exception:
                pass
            r = p.add_run()
            r.text = line
            r.font.name = theme.font_body
            r.font.size = Pt(body_pt)
            r.font.color.rgb = theme.body

    add_page_chrome(
        slide, theme=theme,
        page_n=page_n, page_total=page_total,
        section_label=section_label,
    )
