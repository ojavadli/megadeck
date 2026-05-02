"""Code snippet slide — code block with title + caption."""
from __future__ import annotations

from pptx.enum.text import MSO_ANCHOR
from pptx.slide import Slide
from pptx.util import Inches, Pt

from megadeck.core.schemas import CodeSnippetSlide
from megadeck.design_system.primitives import (
    add_corner_dotgrid,
    add_eyebrow,
    add_oval,
    add_page_chrome,
    add_round_rect,
    add_text,
    set_slide_bg,
)
from megadeck.design_system.tokens import Theme


def render_code_snippet(
    slide: Slide,
    data: CodeSnippetSlide,
    theme: Theme,
    *,
    page_n: int,
    page_total: int,
    section_label: str | None = None,
) -> None:
    set_slide_bg(slide, color=theme.bg, theme=theme)
    add_corner_dotgrid(slide, theme)
    add_eyebrow(slide, text=data.eyebrow.upper(), theme=theme)

    LEFT = theme.left_margin_in
    CONTENT_W = theme.content_width_in

    add_text(
        slide,
        left=LEFT, top=1.20, width=CONTENT_W, height=0.85,
        text=data.title,
        font=theme.font_display,
        size_pt=theme.type_scale.h2,
        color=theme.title,
        bold=True,
    )

    block_top = 2.40
    block_h = 4.10
    # macOS-style window chrome on top of the code block
    chrome_h = 0.30
    add_round_rect(
        slide,
        left=LEFT, top=block_top, width=CONTENT_W, height=chrome_h,
        fill=theme.overlay,
        adjust=0.20,
    )
    # 3 traffic-light dots
    for i, c in enumerate([theme.danger, theme.warning, theme.success]):
        add_oval(
            slide,
            left=LEFT + 0.20 + i * 0.25, top=block_top + 0.07,
            size=0.16, fill=c,
        )
    # Language tag (right-aligned)
    add_text(
        slide,
        left=LEFT + CONTENT_W - 1.2, top=block_top, width=1.10, height=chrome_h,
        text=data.language,
        font=theme.font_mono,
        size_pt=theme.type_scale.eyebrow,
        color=theme.muted,
        bold=True,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    # Code surface
    add_round_rect(
        slide,
        left=LEFT, top=block_top + chrome_h,
        width=CONTENT_W, height=block_h - chrome_h,
        fill=theme.title,  # near-black surface
        line=None,
        adjust=0.04,
    )
    # Code text — monospace, light-on-dark
    code_lines = data.code.replace("\t", "    ")
    tb = slide.shapes.add_textbox(
        Inches(LEFT + 0.30), Inches(block_top + chrome_h + 0.20),
        Inches(CONTENT_W - 0.60), Inches(block_h - chrome_h - 0.40),
    )
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_top = Inches(0.05)
    p = tf.paragraphs[0]
    try:
        p.line_spacing = 1.30
    except Exception:
        pass
    r = p.add_run()
    r.text = code_lines
    r.font.name = theme.font_mono
    r.font.size = Pt(theme.type_scale.micro)
    r.font.color.rgb = theme.inverse  # light on dark

    if data.caption:
        add_text(
            slide,
            left=LEFT, top=block_top + block_h + 0.20,
            width=CONTENT_W, height=0.40,
            text=data.caption,
            font=theme.font_body,
            size_pt=theme.type_scale.micro,
            color=theme.muted,
        )

    add_page_chrome(
        slide, theme=theme,
        page_n=page_n, page_total=page_total,
        section_label=section_label,
    )
